"""
Modul Pembuat Dataset Training Murni (Pure Ground-Truth)
Khusus Mengekstrak dari 4 Berkas Emas: Dataset Asdoku 1 s/d 4.
Standar Acuan: SK Rektor UEU No. 5 Tahun 2018.
"""

import os
import re
import csv
from typing import List, Dict, Any
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

FRONTMATTER_TITLES = [
    "HALAMAN PERNYATAAN KEASLIAN", "PERNYATAAN KEASLIAN",
    "HALAMAN PENGESAHAN TUGAS AKHIR", "HALAMAN PENGESAHAN", "LEMBAR PENGESAHAN",
    "HALAMAN PERSETUJUAN PUBLIKASI KARYA ILMIAH", "PERSETUJUAN PUBLIKASI",
    "KATA PENGANTAR", "PRAKATA", "UCAPAN TERIMA KASIH",
    "ABSTRAK", "ABSTRACT",
    "DAFTAR ISI", "DAFTAR TABEL", "DAFTAR GAMBAR", "DAFTAR SIMBOL", "DAFTAR LAMPIRAN",
    "DAFTAR REFERENSI", "DAFTAR PUSTAKA"
]

CHAPTER_NAMES = [
    "PENDAHULUAN", "TINJAUAN PUSTAKA", "LANDASAN TEORI", "METODE", "METODOLOGI",
    "METODE PENELITIAN", "HASIL DAN PEMBAHASAN", "KESIMPULAN DAN SARAN", "PENUTUP"
]

def clean_text(text: str) -> str:
    if not text:
        return ""
    text = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\x9f]', '', text)
    return re.sub(r'\s+', ' ', text).strip()

def get_font_info(p):
    size, bold, italic = 12.0, False, False
    for r in p.runs:
        if r.text.strip():
            if r.font.size:
                size = r.font.size.pt
            elif p.style and p.style.font and p.style.font.size:
                size = p.style.font.size.pt

            bold = r.font.bold if r.font.bold is not None else (p.style.font.bold if p.style and p.style.font else False)
            italic = r.font.italic if r.font.italic is not None else (p.style.font.italic if p.style and p.style.font else False)
            break
    return size, bool(bold), bool(italic)

def get_alignment_str(align):
    if align == WD_ALIGN_PARAGRAPH.CENTER:
        return "CENTER"
    elif align == WD_ALIGN_PARAGRAPH.JUSTIFY:
        return "JUSTIFY"
    elif align == WD_ALIGN_PARAGRAPH.RIGHT:
        return "RIGHT"
    elif align == WD_ALIGN_PARAGRAPH.LEFT:
        return "LEFT"
    return "DEFAULT"

def extract_and_label_document(docx_path: str, doc_id: str) -> List[Dict[str, Any]]:
    doc = Document(docx_path)
    total_p = len(doc.paragraphs)
    items = []
    current_zone = "COVER"

    for idx, p in enumerate(doc.paragraphs):
        text = clean_text(p.text)
        size, bold, italic = get_font_info(p)
        align = get_alignment_str(p.alignment)
        style_name = p.style.name if p.style else "Normal"
        style_lower = style_name.lower()
        upper = text.upper()
        char_len = len(text)
        word_count = len(text.split())
        doc_pos = idx / float(total_p) if total_p > 0 else 0.0

        pPr = p._p.pPr
        has_num = bool(pPr is not None and pPr.numPr is not None)
        has_img = bool(p._p.xpath(".//w:drawing | .//w:pict"))
        has_pgbrk = bool(p._p.xpath('.//w:br[@w:type="page"]'))

        sb, sa, ls = 0.0, 0.0, 1.15
        fl, li = 0.0, 0.0
        if pPr is not None:
            sp = pPr.find(qn('w:spacing'))
            if sp is not None:
                if qn('w:before') in sp.attrib:
                    sb = float(sp.attrib[qn('w:before')]) / 20.0
                if qn('w:after') in sp.attrib:
                    sa = float(sp.attrib[qn('w:after')]) / 20.0
                if qn('w:line') in sp.attrib:
                    ls = float(sp.attrib[qn('w:line')]) / 240.0
            ind = pPr.find(qn('w:ind'))
            if ind is not None:
                if qn('w:firstLine') in ind.attrib:
                    fl = float(ind.attrib[qn('w:firstLine')]) / 567.0
                elif qn('w:hanging') in ind.attrib:
                    fl = -float(ind.attrib[qn('w:hanging')]) / 567.0
                if qn('w:left') in ind.attrib:
                    li = float(ind.attrib[qn('w:left')]) / 567.0

        upper_ratio = sum(1 for c in text if c.isupper()) / float(char_len) if char_len > 0 else 0.0

        # State tracking zona
        if any(upper == kw for kw in ["HALAMAN PERNYATAAN KEASLIAN", "HALAMAN PENGESAHAN TUGAS AKHIR", "KATA PENGANTAR"]):
            current_zone = "FRONTMATTER"
        elif upper in CHAPTER_NAMES or upper.startswith("BAB "):
            current_zone = "BODY_BAB"
        elif upper in ["DAFTAR REFERENSI", "DAFTAR PUSTAKA"]:
            current_zone = "BIBLIOGRAPHY"
        elif upper.startswith("LAMPIRAN"):
            current_zone = "APPENDIX"

        # Tentukan Label
        label = "BODY_PARAGRAPH"

        if char_len == 0:
            label = "BODY_PARAGRAPH" if has_img else "LAINNYA"

        # 1. Caption Gambar
        elif re.match(r'^(Gambar|Figure)\s+(\d+[\.\-]\d+|\.|\d+).*', text, re.IGNORECASE) and word_count <= 25:
            label = "CAPTION_GAMBAR"

        # 2. Caption Tabel
        elif re.match(r'^(Tabel|Table)\s+(\d+[\.\-]\d+|\.|\d+).*', text, re.IGNORECASE) and word_count <= 25:
            label = "CAPTION_TABEL"

        # 3. Judul Bab Utama (Heading 1 di Zona Bab / Kanonikal)
        elif (
            (re.match(r'^(BAB\s+[IVX\d]+|BAB\s+\d+)', text, re.IGNORECASE) and word_count <= 10)
            or (upper in CHAPTER_NAMES and ("heading 1" in style_lower or size >= 14 or bold))
        ):
            label = "JUDUL_BAB"

        # 4. Sub-Bab (Heading 2 / Heading 3)
        elif (
            "heading 2" in style_lower 
            or "heading 3" in style_lower
            or re.match(r'^\d+\.\d+(\.\d+)?\s+', text)
            or (current_zone == "BODY_BAB" and word_count <= 12 and bold and not text.endswith(".") and size >= 12.0)
        ):
            label = "SUB_BAB"

        # 5. Judul Frontmatter / Halaman Depan / Lampiran
        elif (
            any(upper == kw for kw in FRONTMATTER_TITLES)
            or (upper.startswith("LAMPIRAN") and word_count <= 8)
        ):
            label = "FRONTMATTER_TITLE"

        # 6. Cover Metadata
        elif current_zone == "COVER" and doc_pos < 0.15 and (
            any(k in upper for k in ["TUGAS AKHIR", "SKRIPSI", "ARDITYA ADJIE", "20230801274", "TEKNIK INFORMATIKA", "FAKULTAS ILMU KOMPUTER", "UNIVERSITAS ESA UNGGUL", "TAHUN 2024", "NAMA", "NIM", "DIAJUKAN SEBAGAI"])
        ):
            label = "COVER_METADATA"

        # 7. Entri Daftar Pustaka
        elif current_zone == "BIBLIOGRAPHY" and word_count > 4 and not any(upper == kw for kw in ["DAFTAR PUSTAKA", "DAFTAR REFERENSI"]):
            label = "DAFTAR_PUSTAKA_ITEM"

        # 8. Body Paragraph (Paragraf isi, list rincian, code, math)
        else:
            label = "BODY_PARAGRAPH"

        items.append({
            "doc_id": doc_id,
            "para_id": idx,
            "text": text,
            "style": style_name,
            "char_len": char_len,
            "word_count": word_count,
            "font_size": size,
            "is_bold": bold,
            "is_italic": italic,
            "align": align,
            "has_num": has_num,
            "has_img": has_img,
            "has_pgbrk": has_pgbrk,
            "space_before": round(sb, 2),
            "space_after": round(sa, 2),
            "line_spacing": round(ls, 2),
            "first_line_indent": round(fl, 2),
            "left_indent": round(li, 2),
            "uppercase_ratio": round(upper_ratio, 3),
            "doc_pos": round(doc_pos, 4),
            "label": label
        })

    return items

def main():
    target_files = [
        ("Dataset Asdoku/Dataset Asdoku 1.docx", "ASDOKU_1_TEMPLATE_UTUH"),
        ("Dataset Asdoku/Dataset Asdoku 2.docx", "ASDOKU_2_TOC_HEADINGS"),
        ("Dataset Asdoku/Dataset Asdoku 3.docx", "ASDOKU_3_FIGURES_TOF"),
        ("Dataset Asdoku/Dataset Asdoku 4.docx", "ASDOKU_4_TABLES_TOT")
    ]

    all_samples = []
    for filepath, doc_id in target_files:
        if os.path.exists(filepath):
            print(f"Mengekstrak berkas emas: {filepath}...")
            samples = extract_and_label_document(filepath, doc_id)
            print(f"  -> {len(samples)} baris diekstraksi.")
            all_samples.extend(samples)

    out_csv = "data/dataset_skripsi.csv"
    keys = list(all_samples[0].keys())
    with open(out_csv, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=keys)
        writer.writeheader()
        writer.writerows(all_samples)

    print(f"\n[SUKSES] Dataset training murni berhasil disimpan ke: {out_csv} ({len(all_samples)} baris)")

    # Tampilkan distribusi kelas
    from collections import Counter
    counts = Counter(s["label"] for s in all_samples)
    print("\n--- DISTRIBUSI KELAS TRAINING MURNI ASDOKU ---")
    for k, v in counts.most_common():
        print(f"  {k:22s}: {v:4d} baris")
    print("-----------------------------------------------\n")

if __name__ == "__main__":
    main()
